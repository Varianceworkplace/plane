# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Contract tests for time tracking (plane.app.views.worklog).

Rules under test:
  * time can be logged only when the project has time tracking turned on
  * project admins and members can log time; guests and non-members cannot
  * a worklog can be changed by the member who logged it, a project admin or a workspace admin
  * only workspace admins can list / download all worklogs (Settings > Worklogs)
  * ticket work items stay unreachable
  * start / stop timer: one per member, stop freezes the end time, saving needs a note
"""

import io
from datetime import timedelta
from uuid import uuid4

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from plane.db.models import (
    Issue,
    IssueWorkLog,
    IssueWorkLogTimer,
    Project,
    ProjectMember,
    State,
    User,
    WorkspaceMember,
)


def _user(workspace, first_name, role=15):
    uid = uuid4().hex[:8]
    user = User.objects.create(
        email=f"{first_name.lower()}-{uid}@plane.so",
        username=f"{first_name.lower()}_{uid}",
        first_name=first_name,
        last_name="Test",
        display_name=first_name.lower(),
    )
    WorkspaceMember.objects.create(workspace=workspace, member=user, role=role)
    return user


def _client(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def _project(workspace, owner, name, identifier, enabled=True):
    project = Project.objects.create(
        name=name,
        identifier=identifier,
        workspace=workspace,
        created_by=owner,
        is_time_tracking_enabled=enabled,
    )
    state = State.objects.create(name="Todo", project=project, workspace=workspace, group="backlog", default=True)
    issue = Issue.objects.create(name=f"{name} task", workspace=workspace, project=project, state=state)
    return project, issue


@pytest.fixture
def setup(db, workspace, create_user):
    people = {
        "admin": create_user,  # workspace admin
        "member": _user(workspace, "Hiral"),
        "other": _user(workspace, "Vijay"),
        "guest": _user(workspace, "Guest", role=5),
        "outsider": _user(workspace, "Outsider"),
    }
    project, issue = _project(workspace, create_user, "Power Layout", "POWER")
    off_project, off_issue = _project(workspace, create_user, "No Tracking", "NOTRK", enabled=False)
    for proj in (project, off_project):
        ProjectMember.objects.create(project=proj, member=people["admin"], role=20, is_active=True)
        ProjectMember.objects.create(project=proj, member=people["member"], role=15, is_active=True)
        ProjectMember.objects.create(project=proj, member=people["other"], role=15, is_active=True)
        ProjectMember.objects.create(project=proj, member=people["guest"], role=5, is_active=True)
    return {
        "workspace": workspace,
        "people": people,
        "clients": {name: _client(user) for name, user in people.items()},
        "project": project,
        "issue": issue,
        "off_project": off_project,
        "off_issue": off_issue,
    }


def _url(setup, project=None, issue=None, worklog_id=None):
    project = project or setup["project"]
    issue = issue or setup["issue"]
    url = f"/api/workspaces/{setup['workspace'].slug}/projects/{project.id}/issues/{issue.id}/worklogs/"
    return f"{url}{worklog_id}/" if worklog_id else url


def _log(setup, who="member", duration=90, description="create zip"):
    response = setup["clients"][who].post(
        _url(setup), {"duration": duration, "description": description}, format="json"
    )
    assert response.status_code == status.HTTP_201_CREATED, response.content
    return response.json()


# ---------------------------------------------------------------- logging time


@pytest.mark.contract
def test_member_logs_time(setup):
    worklog = _log(setup)
    assert worklog["duration"] == 90
    assert worklog["description"] == "create zip"
    assert worklog["logged_by"]["id"] == str(setup["people"]["member"].id)
    assert worklog["issue"]["id"] == str(setup["issue"].id)
    assert worklog["project"]["identifier"] == "POWER"

    response = setup["clients"]["guest"].get(_url(setup))
    assert response.status_code == status.HTTP_200_OK
    assert [w["id"] for w in response.json()] == [worklog["id"]]


@pytest.mark.contract
def test_cannot_log_when_time_tracking_is_off(setup):
    response = setup["clients"]["member"].post(
        _url(setup, setup["off_project"], setup["off_issue"]), {"duration": 30}, format="json"
    )
    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert not IssueWorkLog.objects.filter(issue=setup["off_issue"]).exists()


@pytest.mark.contract
@pytest.mark.parametrize("duration", [0, -5, "abc", None, 999 * 60 + 60])
def test_invalid_duration_is_rejected(setup, duration):
    response = setup["clients"]["member"].post(_url(setup), {"duration": duration}, format="json")
    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert "duration" in response.json()


@pytest.mark.contract
@pytest.mark.parametrize("who", ["guest", "outsider"])
def test_guests_and_non_members_cannot_log(setup, who):
    response = setup["clients"][who].post(_url(setup), {"duration": 30}, format="json")
    assert response.status_code == status.HTTP_403_FORBIDDEN


# ---------------------------------------------------------------- edit / delete


@pytest.mark.contract
def test_owner_can_edit_and_delete_own_worklog(setup):
    worklog = _log(setup)
    client = setup["clients"]["member"]
    response = client.patch(_url(setup, worklog_id=worklog["id"]), {"duration": 135}, format="json")
    assert response.status_code == status.HTTP_200_OK
    assert response.json()["duration"] == 135
    assert response.json()["description"] == "create zip"

    response = client.delete(_url(setup, worklog_id=worklog["id"]))
    assert response.status_code == status.HTTP_204_NO_CONTENT
    assert client.get(_url(setup)).json() == []


@pytest.mark.contract
def test_other_member_cannot_change_someone_elses_worklog(setup):
    worklog = _log(setup)
    client = setup["clients"]["other"]
    assert client.patch(_url(setup, worklog_id=worklog["id"]), {"duration": 5}, format="json").status_code == 403
    assert client.delete(_url(setup, worklog_id=worklog["id"])).status_code == 403
    assert IssueWorkLog.objects.get(pk=worklog["id"]).duration == 90


@pytest.mark.contract
def test_admin_can_change_any_worklog(setup):
    worklog = _log(setup)
    client = setup["clients"]["admin"]
    assert client.patch(_url(setup, worklog_id=worklog["id"]), {"duration": 60}, format="json").status_code == 200
    assert client.delete(_url(setup, worklog_id=worklog["id"])).status_code == 204


# ---------------------------------------------------------------- workspace worklogs (settings)


def _ws_url(setup, suffix=""):
    return f"/api/workspaces/{setup['workspace'].slug}/worklogs/{suffix}"


@pytest.mark.contract
def test_only_workspace_admin_sees_all_worklogs(setup):
    _log(setup)
    assert setup["clients"]["member"].get(_ws_url(setup)).status_code == status.HTTP_403_FORBIDDEN
    assert setup["clients"]["member"].get(_ws_url(setup, "download/")).status_code == status.HTTP_403_FORBIDDEN

    response = setup["clients"]["admin"].get(_ws_url(setup))
    assert response.status_code == status.HTTP_200_OK
    body = response.json()
    assert body["total_count"] == 1
    assert body["total_duration"] == 90


@pytest.mark.contract
def test_workspace_worklog_filters(setup):
    mine = _log(setup, "member", 90)
    _log(setup, "other", 30)
    old = IssueWorkLog.objects.get(pk=mine["id"])
    old.created_at = timezone.now() - timedelta(days=40)
    old.save(update_fields=["created_at"])

    client = setup["clients"]["admin"]
    by_user = client.get(_ws_url(setup), {"logged_by": str(setup["people"]["member"].id)}).json()
    assert [w["id"] for w in by_user["results"]] == [mine["id"]]

    start = (timezone.now() - timedelta(days=7)).isoformat()
    recent = client.get(_ws_url(setup), {"start": start}).json()
    assert recent["total_count"] == 1
    assert recent["total_duration"] == 30

    other_project = client.get(_ws_url(setup), {"project": str(setup["off_project"].id)}).json()
    assert other_project["total_count"] == 0


@pytest.mark.contract
def test_download_csv_and_xlsx(setup):
    _log(setup, "member", 630, "=HYPERLINK(evil)")
    client = setup["clients"]["admin"]

    response = client.get(_ws_url(setup, "download/"), {"type": "csv"})
    assert response.status_code == status.HTTP_200_OK
    assert response["Content-Type"].startswith("text/csv")
    text = response.content.decode()
    assert "Project,Work item ID,Work item,Logged by,Email,Date,Start time,End time,Time" in text
    assert "POWER-" in text and "10h 30m" in text and "10.5" in text
    assert "'=HYPERLINK(evil)" in text  # formula injection neutralised

    response = client.get(_ws_url(setup, "download/"), {"type": "xlsx"})
    assert response.status_code == status.HTTP_200_OK
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Project"
    assert rows[1][0] == "Power Layout" and rows[1][9] == 630
    assert rows[1][6] in ("", None) and rows[1][7] in ("", None)  # manual entry: no start / end

    assert client.get(_ws_url(setup, "download/"), {"type": "pdf"}).status_code == 400


# ---------------------------------------------------------------- project flag and ticket isolation


@pytest.mark.contract
def test_project_list_exposes_time_tracking_flag(setup):
    response = setup["clients"]["admin"].get(f"/api/workspaces/{setup['workspace'].slug}/projects/")
    assert response.status_code == status.HTTP_200_OK
    projects = {p["identifier"]: p for p in response.json()}
    assert projects["POWER"]["is_time_tracking_enabled"] is True
    assert projects["NOTRK"]["is_time_tracking_enabled"] is False


@pytest.mark.contract
def test_ticket_work_items_are_unreachable(setup):
    slug = setup["workspace"].slug
    client = setup["clients"]["member"]
    meta = client.get(f"/api/workspaces/{slug}/tickets/meta/").json()
    ticket = client.post(
        f"/api/workspaces/{slug}/tickets/",
        {
            "name": "Ticket",
            "assignee_id": str(setup["people"]["member"].id),
            "manager_id": str(setup["people"]["member"].id),
            "state_id": meta["statuses"][0]["id"],
        },
        format="json",
    )
    assert ticket.status_code == status.HTTP_201_CREATED, ticket.content
    ticket_issue = Issue.all_objects.get(pk=ticket.json()["issue_id"])
    ticket_project = Project.all_objects.get(pk=ticket_issue.project_id)
    ticket_project.is_time_tracking_enabled = True
    ticket_project.save(update_fields=["is_time_tracking_enabled"])

    url = _url(setup, ticket_project, ticket_issue)
    assert client.post(url, {"duration": 30}, format="json").status_code == status.HTTP_404_NOT_FOUND
    assert client.get(url).status_code == status.HTTP_404_NOT_FOUND
    assert not IssueWorkLog.objects.filter(issue_id=ticket_issue.id).exists()


# ---------------------------------------------------------------- start / stop timer


def _timer_url(setup, suffix=""):
    return f"/api/workspaces/{setup['workspace'].slug}/worklog-timer/{suffix}"


def _start_url(setup, project=None, issue=None):
    project = project or setup["project"]
    issue = issue or setup["issue"]
    return f"/api/workspaces/{setup['workspace'].slug}/projects/{project.id}/issues/{issue.id}/worklog-timer/"


@pytest.mark.contract
def test_timer_start_stop_save_creates_worklog_with_times(setup):
    client = setup["clients"]["member"]
    assert client.get(_timer_url(setup)).json() == {"timer": None}

    response = client.post(_start_url(setup))
    assert response.status_code == status.HTTP_201_CREATED, response.content
    assert response.json()["is_running"] is True
    assert response.json()["issue"]["id"] == str(setup["issue"].id)

    # pretend the member worked 1h 25m 40s
    timer = IssueWorkLogTimer.objects.get(user=setup["people"]["member"])
    timer.started_at = timezone.now() - timedelta(hours=1, minutes=25, seconds=40)
    timer.save(update_fields=["started_at"])

    stopped = client.post(_timer_url(setup, "stop/"))
    assert stopped.status_code == status.HTTP_200_OK
    body = stopped.json()
    assert body["is_running"] is False and body["stopped_at"]
    assert body["duration"] == 86  # rounded to the nearest minute

    # stopping again keeps the first end time
    assert client.post(_timer_url(setup, "stop/")).json()["stopped_at"] == body["stopped_at"]

    assert client.post(_timer_url(setup, "save/"), {"description": "  "}, format="json").status_code == 400

    saved = client.post(_timer_url(setup, "save/"), {"description": "Fixed language files"}, format="json")
    assert saved.status_code == status.HTTP_201_CREATED, saved.content
    worklog = saved.json()
    assert worklog["duration"] == 86
    assert worklog["description"] == "Fixed language files"
    assert worklog["started_at"] and worklog["ended_at"]
    assert client.get(_timer_url(setup)).json() == {"timer": None}

    stored = IssueWorkLog.objects.get(pk=worklog["id"])
    # the end time is the moment Stop was pressed, not when the note was saved
    assert abs((stored.ended_at - stored.started_at).total_seconds() - (85 * 60 + 40)) < 5


@pytest.mark.contract
def test_only_one_timer_per_member(setup):
    client = setup["clients"]["member"]
    assert client.post(_start_url(setup)).status_code == 201
    second_issue = Issue.objects.create(
        name="Second", workspace=setup["workspace"], project=setup["project"], state=setup["issue"].state
    )
    response = client.post(_start_url(setup, issue=second_issue))
    assert response.status_code == status.HTTP_409_CONFLICT
    assert response.json()["timer"]["issue"]["id"] == str(setup["issue"].id)

    # someone else can still run their own timer
    assert setup["clients"]["other"].post(_start_url(setup, issue=second_issue)).status_code == 201


@pytest.mark.contract
def test_discard_timer_logs_nothing(setup):
    client = setup["clients"]["member"]
    client.post(_start_url(setup))
    assert client.delete(_timer_url(setup)).status_code == status.HTTP_204_NO_CONTENT
    assert client.get(_timer_url(setup)).json() == {"timer": None}
    assert not IssueWorkLog.objects.filter(issue=setup["issue"]).exists()
    # and a new timer can be started again
    assert client.post(_start_url(setup)).status_code == 201


@pytest.mark.contract
def test_timer_rules(setup):
    # guests and projects without time tracking cannot start a timer
    assert setup["clients"]["guest"].post(_start_url(setup)).status_code == status.HTTP_403_FORBIDDEN
    response = setup["clients"]["member"].post(_start_url(setup, setup["off_project"], setup["off_issue"]))
    assert response.status_code == status.HTTP_400_BAD_REQUEST
    # nothing to stop or save
    assert setup["clients"]["member"].post(_timer_url(setup, "stop/")).status_code == 404
    assert (
        setup["clients"]["member"].post(_timer_url(setup, "save/"), {"description": "x"}, format="json").status_code
        == 404
    )


@pytest.mark.contract
def test_editing_timer_entry_moves_end_time(setup):
    client = setup["clients"]["member"]
    client.post(_start_url(setup))
    worklog = client.post(_timer_url(setup, "save/"), {"description": "work"}, format="json").json()
    response = client.patch(_url(setup, worklog_id=worklog["id"]), {"duration": 120}, format="json")
    assert response.status_code == 200
    stored = IssueWorkLog.objects.get(pk=worklog["id"])
    assert stored.ended_at - stored.started_at == timedelta(minutes=120)
